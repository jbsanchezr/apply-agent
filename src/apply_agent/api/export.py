"""The applications table as an Excel workbook.

Every text cell comes from email content or the LLM, so it is written as an
explicit string. Left to itself, openpyxl stores a value starting with ``=``
as a formula, which would let a hostile email run one in the user's Excel.

Colours match the page: the status cell carries its badge colour, the stage
cell a soft tint, and an open application with no news lately has its day
count in the same amber.
"""

from collections.abc import Sequence
from datetime import datetime
from io import BytesIO
from typing import Final

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from apply_agent.api.page import QUIET_COLOUR, STATUS_COLOURS
from apply_agent.domain import ApplicationStage, stage_of
from apply_agent.storage.repository import ApplicationView

XLSX_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS: Final = (
    "Company",
    "Role",
    "Stage",
    "Status",
    "Set by hand",
    "Latest",
    "First seen",
    "Last news",
    "Days since last news",
)
_WIDTHS: Final = (28, 36, 12, 22, 12, 70, 12, 12, 12)
_DATE_FORMAT: Final = "yyyy-mm-dd"
_STAGE_TINTS: Final = {
    ApplicationStage.ADVANCING: "#dcfce7",
    ApplicationStage.SENT: "#f1f5f9",
    ApplicationStage.REJECTED: "#fee2e2",
}
_WHITE: Final = "FFFFFFFF"
_DARK: Final = "FF1F2937"


def _font(*, color: str | None = None, bold: bool = False) -> Font:
    # Name the font, or styled cells fall back to the viewer's default face.
    return Font(name="Calibri", size=11, color=color, bold=bold)


def _argb(colour: str) -> str:
    return "FF" + colour.lstrip("#").upper()


def _fill(colour: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=_argb(colour), end_color=_argb(colour))


def _text(sheet: Worksheet, row: int, column: int, value: str) -> Cell:
    cell = sheet.cell(row=row, column=column, value=value)
    cell.data_type = "s"
    return cell


def applications_workbook(
    views: Sequence[ApplicationView], now: datetime, *, quiet_after_days: int
) -> bytes:
    workbook = Workbook(write_only=False)
    workbook.remove(workbook.worksheets[0])
    sheet = workbook.create_sheet("Applications")
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = _font(bold=True)

    for row, view in enumerate(views, start=2):
        app = view.application
        stage = stage_of(app.status)
        texts = (
            app.company,
            app.role or "",
            stage.value,
            app.status.value.replace("_", " "),
            "yes" if app.status_overridden else "",
            view.latest_summary,
        )
        cells = [_text(sheet, row, column, value) for column, value in enumerate(texts, start=1)]
        cells[2].fill = _fill(_STAGE_TINTS[stage])
        cells[2].font = _font(color=_DARK)
        cells[3].fill = _fill(STATUS_COLOURS[app.status.value])
        cells[3].font = _font(color=_WHITE, bold=True)
        for column, when in ((7, app.first_seen_at), (8, app.last_activity_at)):
            cell = sheet.cell(row=row, column=column, value=when.date())
            cell.number_format = _DATE_FORMAT
        days = max((now - app.last_activity_at).days, 0)
        waiting = sheet.cell(row=row, column=9, value=days)
        if stage is not ApplicationStage.REJECTED and days >= quiet_after_days:
            waiting.font = _font(color=_argb(QUIET_COLOUR), bold=True)

    for column, width in enumerate(_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    # Printed or exported to PDF, the table fits the page width instead of
    # losing its right-hand columns to a second page.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_title_rows = "1:1"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

"""The HTTP API, end to end through FastAPI's test client."""

import base64
import json
import re
import threading
import time
from collections.abc import Iterator, Sequence
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
import pytest
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage
from openpyxl import load_workbook
from sqlalchemy import Engine

from apply_agent.agent.outcomes import SyncReport
from apply_agent.api.app import create_app
from apply_agent.api.export import HEADERS, XLSX_MEDIA_TYPE
from apply_agent.api.page import gmail_url
from apply_agent.config import Settings
from apply_agent.domain import MessageAssessment
from apply_agent.domain.offers import JobOffer, OfferSource
from apply_agent.domain.profile import CandidateProfile
from apply_agent.domain.search import BoardStatus
from apply_agent.offers.boards import CompanyBoards
from apply_agent.offers.company import CompanyResearcher, NewsItem
from apply_agent.offers.infojobs import InfoJobsSearch
from apply_agent.offers.scanner import OfferScanReport
from apply_agent.offers.search import JobSearcher
from apply_agent.storage import make_session_factory
from apply_agent.storage.offers import OfferRepository
from apply_agent.storage.repository import Repository
from apply_agent.storage.searches import SearchRepository
from apply_agent.storage.targets import TargetRepository
from tests.agent.helpers import T0, ScriptedChatModel, assessment, message
from tests.conftest import FIXTURES_DIR
from tests.offers.test_boards import GREENHOUSE
from tests.offers.test_infojobs import SECRET as INFOJOBS_SECRET
from tests.offers.test_infojobs import item as infojobs_item
from tests.offers.test_web import investigator as web_investigator

GOOGLE_NEWS = (
    '<?xml version="1.0"?><rss><channel><item><title>Arcwell raises EUR 40M - Wire</title>'
    "<link>https://news.google.com/rss/articles/a</link>"
    '<pubDate>Wed, 23 Sep 2026 07:00:00 GMT</pubDate><source url="https://w.example">Wire</source>'
    "</item></channel></rss>"
)

TOKEN = "test-token-not-secret"  # noqa: S105 - fake token for tests
AUTH = {"Authorization": f"Bearer {TOKEN}"}
EMPTY_REPORT = SyncReport(processed=0, recorded=0, failed=0, by_category={}, failures=[])
# The day after the bundled LinkedIn alert, so its offers are not old yet.
ALERT_DAY = datetime(2026, 9, 25, 9, 0, tzinfo=UTC)


def _settings(tmp_path: Path, **extra: str) -> Settings:
    return Settings.from_env(
        {
            "APPLY_AGENT_DATABASE_URL": f"sqlite:///{tmp_path / 'api.db'}",
            "APPLY_AGENT_FIXTURES_DIR": str(FIXTURES_DIR / "emails"),
            "APPLY_AGENT_INITIAL_LOOKBACK_DAYS": "3650",
            "APPLY_AGENT_API_TOKEN": TOKEN,
            # Never the real profile in ~/.config: tests must not depend on it.
            "APPLY_AGENT_PROFILE_PATH": str(tmp_path / "profile.toml"),
            **extra,
        }
    )


def _wait_for(client: TestClient, job_id: str, timeout: float = 30.0) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job: dict[str, Any] = client.get(f"/sync/{job_id}", headers=AUTH).json()
        if job["status"] != "running":
            return job
        time.sleep(0.05)
    raise AssertionError(f"sync {job_id} did not finish within {timeout}s")


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    """The real app: real agent (keyword baseline) over the fixtures, real SQLite."""
    with TestClient(create_app(_settings(tmp_path))) as test_client:
        yield test_client


@pytest.fixture
def repository(engine: Engine) -> Repository:
    return Repository(make_session_factory(engine))


def test_health_and_metrics_need_no_token(client: TestClient) -> None:
    assert client.get("/healthz").json() == {"status": "ok"}
    body = client.get("/metrics").text
    assert "# TYPE apply_agent_sync_runs_total counter" in body
    assert "apply_agent_applications" in body


def test_missing_token_gets_a_browser_login_prompt(client: TestClient) -> None:
    response = client.get("/applications")
    assert response.status_code == 401
    assert response.headers["www-authenticate"].startswith("Basic")


@pytest.mark.parametrize(
    "header",
    [
        f"Bearer {TOKEN}",
        "Basic " + base64.b64encode(f"anyone:{TOKEN}".encode()).decode(),
    ],
)
def test_bearer_and_basic_are_both_accepted(client: TestClient, header: str) -> None:
    assert client.get("/applications", headers={"Authorization": header}).status_code == 200


@pytest.mark.parametrize(
    "header",
    [
        "Bearer wrong",
        "Basic !!!not-base64",
        "Basic " + base64.b64encode(b"user:").decode(),
        "Token x",
    ],
)
def test_bad_credentials_are_rejected(client: TestClient, header: str) -> None:
    assert client.get("/applications", headers={"Authorization": header}).status_code == 401


def test_sync_runs_in_the_background_and_fills_the_table(client: TestClient) -> None:
    started = client.post("/sync", headers=AUTH)
    assert started.status_code == 202

    job = _wait_for(client, started.json()["id"])
    assert job["status"] == "succeeded"
    assert job["report"]["processed"] == len(list((FIXTURES_DIR / "emails").glob("*.eml")))

    rows = client.get("/applications?format=json", headers=AUTH).json()
    assert rows
    assert {"company", "status", "latest_summary", "thread_ids"} <= set(rows[0])
    page = client.get("/applications", headers=AUTH).text
    assert "Kestrel Robotics" in page
    assert "Last sync: succeeded" in page

    again = _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
    assert again["report"]["processed"] == 0, "a second sync must be idempotent"


def test_json_is_also_chosen_by_accept_header(client: TestClient) -> None:
    response = client.get("/applications", headers={**AUTH, "Accept": "application/json"})
    assert response.headers["content-type"].startswith("application/json")


def test_a_second_sync_is_refused_while_one_is_running(
    tmp_path: Path, repository: Repository
) -> None:
    release = threading.Event()

    def slow_sync() -> SyncReport:
        release.wait(timeout=10)
        return EMPTY_REPORT

    app = create_app(_settings(tmp_path), run_sync=slow_sync, repository=repository)
    with TestClient(app) as client:
        first = client.post("/sync", headers=AUTH)
        second = client.post("/sync", headers=AUTH)
        release.set()

        assert first.status_code == 202
        assert second.status_code == 409
        assert second.json()["running"]["id"] == first.json()["id"]
        assert _wait_for(client, first.json()["id"])["status"] == "succeeded"
        assert client.post("/sync", headers=AUTH).status_code == 202, "the lock is released"


def test_a_failing_sync_is_reported_and_does_not_block_the_next(
    tmp_path: Path, repository: Repository
) -> None:
    def broken_sync() -> SyncReport:
        raise ConnectionError("mail server unreachable")

    app = create_app(_settings(tmp_path), run_sync=broken_sync, repository=repository)
    with TestClient(app) as client:
        job = _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
        assert job["status"] == "failed"
        assert "mail server unreachable" in job["error"]
        assert client.post("/sync", headers=AUTH).status_code == 202


def test_unknown_job_is_404(client: TestClient) -> None:
    assert client.get("/sync/does-not-exist", headers=AUTH).status_code == 404


def test_email_derived_text_is_escaped_in_the_page(tmp_path: Path, repository: Repository) -> None:
    hostile = MessageAssessment.model_validate(
        assessment("offer", company="<script>alert(1)</script>", summary="<img src=x onerror=1>")
    )
    repository.record(message("m1"), hostile)
    app = create_app(_settings(tmp_path), run_sync=lambda: EMPTY_REPORT, repository=repository)
    with TestClient(app) as client:
        page = client.get("/applications", headers=AUTH).text
    assert "<script>alert(1)</script>" not in page
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in page
    assert "<img src=x" not in page


def test_without_a_configured_token_one_is_generated_and_printed(
    tmp_path: Path, repository: Repository, capsys: pytest.CaptureFixture[str]
) -> None:
    settings = _settings(tmp_path).model_copy(update={"api_token": None})
    app = create_app(settings, run_sync=lambda: EMPTY_REPORT, repository=repository)
    printed = re.search(r"^\s{4}(\S+)$", capsys.readouterr().err, re.MULTILINE)
    assert printed is not None
    with TestClient(app) as client:
        assert client.get("/applications").status_code == 401
        headers = {"Authorization": f"Bearer {printed.group(1)}"}
        assert client.get("/applications", headers=headers).status_code == 200


def _section(page: str, stage: str) -> str:
    match = re.search(rf'<section id="{stage}">(.*?)</section>', page, re.S)
    assert match, f"no {stage} section"
    return match.group(1)


def test_applications_are_grouped_by_stage_with_quiet_ones_flagged(
    tmp_path: Path, repository: Repository
) -> None:
    """Sent day 0, Nubaria interviewed day 2, Arcwell rejected day 5; viewed on day 30."""
    sent = assessment("other", company="Lumora Energy", role="Data Analyst")
    advancing = assessment("interview_invitation", company="Nubaria", role="ML Engineer")
    rejected = assessment("rejection", company="Arcwell Systems", role="Data Engineer")
    for i, (raw, days) in enumerate([(sent, 0), (advancing, 2), (rejected, 5)]):
        repository.record(
            message(f"m{i}", f"t{i}", days=days), MessageAssessment.model_validate(raw)
        )

    app = create_app(
        _settings(tmp_path, APPLY_AGENT_QUIET_AFTER_DAYS="28"),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: T0 + timedelta(days=30),
    )
    with TestClient(app) as client:
        page = client.get("/applications", headers=AUTH).text
        rows = client.get("/applications?format=json", headers=AUTH).json()

    assert "Nubaria" in _section(page, "advancing")
    assert "Lumora Energy" in _section(page, "sent")
    assert "Arcwell Systems" in _section(page, "rejected")
    assert page.index('id="advancing"') < page.index('id="sent"') < page.index('id="rejected"')
    assert "<details>" in _section(page, "rejected"), "rejections are folded away"

    # 30 days without news on the sent one (quiet), 28 on the advancing one
    # (quiet, at the threshold), 25 on the rejection (never flagged).
    assert '<span class="quiet">30 days ago</span>' in _section(page, "sent")
    assert '<span class="quiet">28 days ago</span>' in _section(page, "advancing")
    assert '<span class="muted">25 days ago</span>' in _section(page, "rejected")
    assert "2 with no news lately" in page

    by_company = {r["company"]: r for r in rows}
    assert by_company["Lumora Energy"]["stage"] == "sent"
    assert by_company["Nubaria"]["stage"] == "advancing"
    assert by_company["Arcwell Systems"]["stage"] == "rejected"
    assert by_company["Lumora Energy"]["days_since_last_news"] == 30


def test_an_empty_table_asks_for_a_sync(tmp_path: Path, repository: Repository) -> None:
    app = create_app(_settings(tmp_path), run_sync=lambda: EMPTY_REPORT, repository=repository)
    with TestClient(app) as client:
        page = client.get("/applications", headers=AUTH).text
    assert "No applications yet. Run a sync." in page
    assert "<section" not in page


@pytest.fixture
def seeded(tmp_path: Path, repository: Repository) -> Iterator[TestClient]:
    """One application wrongly marked as advancing, seen 10 days later."""
    wrong = assessment("information_request", company="Halcyon Health", role="MLOps Engineer")
    repository.record(message("m1"), MessageAssessment.model_validate(wrong))
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: T0 + timedelta(days=10),
    )
    with TestClient(app) as client:
        yield client


def _first(client: TestClient) -> dict[str, Any]:
    row: dict[str, Any] = client.get("/applications?format=json", headers=AUTH).json()[0]
    return row


def test_a_false_positive_can_be_corrected_and_the_correction_undone(seeded: TestClient) -> None:
    app_id = _first(seeded)["id"]
    assert "Halcyon Health" in _section(seeded.get("/applications", headers=AUTH).text, "advancing")

    fixed = seeded.put(f"/applications/{app_id}/status", headers=AUTH, json={"status": "applied"})

    assert fixed.status_code == 200
    assert fixed.json()["status"] == "applied"
    assert fixed.json()["stage"] == "sent"
    assert _first(seeded)["status_overridden"] is True
    page = seeded.get("/applications", headers=AUTH).text
    assert "Halcyon Health" in _section(page, "sent")
    assert "set by hand" in _section(page, "sent")
    assert '<option value="auto">' in _section(page, "sent")

    undone = seeded.delete(f"/applications/{app_id}/status", headers=AUTH)

    assert undone.json()["status"] == "information_requested"
    assert _first(seeded)["status_overridden"] is False
    assert "set by hand" not in seeded.get("/applications", headers=AUTH).text


def test_the_correction_menu_offers_every_other_status(seeded: TestClient) -> None:
    page = seeded.get("/applications", headers=AUTH).text
    menu = re.search(r"<select.*?</select>", page, re.S)
    assert menu
    values = re.findall(r'<option value="([^"]*)"', menu.group(0))
    assert values == ["", "applied", "interviewing", "offer_received", "rejected"]


@pytest.mark.parametrize(
    ("method", "path", "body", "expected"),
    [
        ("put", "/applications/999/status", {"status": "applied"}, 404),
        ("delete", "/applications/999/status", None, 404),
        ("put", "/applications/1/status", {"status": "hired"}, 422),
        ("put", "/applications/1/status", {"status": "applied", "extra": 1}, 422),
    ],
)
def test_bad_corrections_are_rejected(
    seeded: TestClient, method: str, path: str, body: dict[str, Any] | None, expected: int
) -> None:
    response = seeded.request(method, path, headers=AUTH, json=body)
    assert response.status_code == expected
    assert _first(seeded)["status"] == "information_requested"


@pytest.mark.parametrize("method", ["put", "delete"])
def test_corrections_need_the_token(seeded: TestClient, method: str) -> None:
    response = seeded.request(method, "/applications/1/status", json={"status": "applied"})
    assert response.status_code == 401


def _sheet(client: TestClient) -> Any:
    response = client.get("/applications?format=xlsx", headers=AUTH)
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="applications-2026-07-11.xlsx"'
    )
    return load_workbook(BytesIO(response.content))["Applications"]


def test_excel_export_has_one_row_per_application(seeded: TestClient) -> None:
    seeded.put("/applications/1/status", headers=AUTH, json={"status": "rejected"})

    rows = list(_sheet(seeded).iter_rows(values_only=True))

    assert rows[0] == HEADERS
    *texts, first_seen, last_news, days = rows[1]
    assert texts == [
        "Halcyon Health",
        "MLOps Engineer",
        "rejected",
        "rejected",
        "yes",
        "Something happened",
    ]
    assert first_seen.date() == last_news.date() == date(2026, 7, 1)
    assert days == 10
    assert len(rows) == 2


def test_excel_export_never_writes_email_text_as_a_formula(
    tmp_path: Path, repository: Repository
) -> None:
    """An email can put anything in a company name; Excel must not run it."""
    hostile = assessment(
        "offer", company='=HYPERLINK("http://evil.example","x")', role="+1+1", summary="@SUM(A1)"
    )
    repository.record(message("m1"), MessageAssessment.model_validate(hostile))
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: T0 + timedelta(days=10),
    )
    with TestClient(app) as client:
        sheet = _sheet(client)

    company, role, *_, latest = (sheet.cell(row=2, column=c) for c in (1, 2, 6))
    assert company.data_type == role.data_type == latest.data_type == "s"
    assert company.value == '=HYPERLINK("http://evil.example","x")'


def test_the_page_has_a_search_box_over_company_and_role(seeded: TestClient) -> None:
    page = seeded.get("/applications", headers=AUTH).text
    assert '<input type="search" id="search"' in page
    assert '<tr data-search="Halcyon Health MLOps Engineer">' in page


def test_the_page_links_to_the_excel_export(seeded: TestClient) -> None:
    assert 'href="applications?format=xlsx"' in seeded.get("/applications", headers=AUTH).text


def test_excel_export_is_colour_coded_like_the_page(tmp_path: Path, repository: Repository) -> None:
    """Sent 30 days ago (quiet), interviewing 2 days ago, rejected 25 days ago."""
    rows = [
        (assessment("other", company="Lumora Energy", role="Data Analyst"), 0),
        (assessment("interview_invitation", company="Nubaria", role="ML Engineer"), 28),
        (assessment("rejection", company="Arcwell Systems", role="Data Engineer"), 5),
    ]
    for i, (raw, days) in enumerate(rows):
        repository.record(
            message(f"m{i}", f"t{i}", days=days), MessageAssessment.model_validate(raw)
        )
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: T0 + timedelta(days=30),
    )
    with TestClient(app) as client:
        response = client.get("/applications?format=xlsx", headers=AUTH)
    sheet = load_workbook(BytesIO(response.content))["Applications"]
    by_company = {row[0].value: row for row in sheet.iter_rows(min_row=2)}

    def colours(company: str) -> tuple[str, str, str | None]:
        _, _, stage, status, *_, days = by_company[company]
        font = days.font.color
        return stage.fill.start_color.rgb, status.fill.start_color.rgb, font and font.rgb

    assert colours("Lumora Energy") == ("FFF1F5F9", "FF64748B", "FFB45309")  # sent, quiet
    assert colours("Nubaria")[:2] == ("FFDCFCE7", "FF1D4ED8")  # advancing, interviewing
    assert colours("Nubaria")[2] != "FFB45309", "2 days without news is not quiet"
    assert colours("Arcwell Systems")[:2] == ("FFFEE2E2", "FFB91C1C")  # rejected
    assert colours("Arcwell Systems")[2] != "FFB45309", "a rejection is never flagged as quiet"
    assert by_company["Nubaria"][3].font.color.rgb == "FFFFFFFF"


@pytest.fixture
def duplicated(tmp_path: Path, repository: Repository) -> Iterator[TestClient]:
    """One application under two spellings, another role there, and an unrelated company."""
    rows = [
        assessment("interview_invitation", company="Cinderpeak Games", role="Backend Engineer"),
        assessment("offer", company="Cinderpeak", role="Backend Engineer - Offer"),
        assessment("other", company="Cinderpeak Games", role="Data Engineer"),
        assessment("other", company="Nubaria", role="ML Engineer"),
    ]
    for i, raw in enumerate(rows):
        repository.record(message(f"m{i}", f"t{i}", days=i), MessageAssessment.model_validate(raw))
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: T0 + timedelta(days=10),
    )
    with TestClient(app) as client:
        yield client


def _ids(client: TestClient) -> dict[tuple[str, str], int]:
    rows = client.get("/applications?format=json", headers=AUTH).json()
    return {(r["company"], r["role"]): r["id"] for r in rows}


def _menu(page: str, application_id: int) -> str:
    match = re.search(
        rf'<select[^>]*onchange="correct\({application_id}, this\)">.*?</select>', page, re.S
    )
    assert match, f"no menu for application {application_id}"
    return match.group(0)


def test_the_menu_offers_merging_only_with_similar_companies(duplicated: TestClient) -> None:
    ids = _ids(duplicated)
    page = duplicated.get("/applications", headers=AUTH).text
    duplicate = ids["Cinderpeak", "Backend Engineer - Offer"]
    original = ids["Cinderpeak Games", "Backend Engineer"]
    other_role = ids["Cinderpeak Games", "Data Engineer"]

    offered = re.findall(r'value="merge:(\d+)"', _menu(page, duplicate))

    assert sorted(map(int, offered)) == sorted([original, other_role])
    assert "Backend Engineer (Cinderpeak Games)" in _menu(page, duplicate)
    assert "merge:" not in _menu(page, ids["Nubaria", "ML Engineer"])


def test_a_duplicate_can_be_merged_into_the_real_application(duplicated: TestClient) -> None:
    ids = _ids(duplicated)
    duplicate = ids["Cinderpeak", "Backend Engineer - Offer"]
    original = ids["Cinderpeak Games", "Backend Engineer"]

    merged = duplicated.post(
        f"/applications/{duplicate}/merge", headers=AUTH, json={"into": original}
    )

    assert merged.status_code == 200
    assert merged.json()["id"] == original
    assert merged.json()["status"] == "offer_received"
    assert set(_ids(duplicated)) == {
        ("Cinderpeak Games", "Backend Engineer"),
        ("Cinderpeak Games", "Data Engineer"),
        ("Nubaria", "ML Engineer"),
    }


@pytest.mark.parametrize(
    ("source", "body", "expected"),
    [
        (1, {"into": 1}, 422),
        (1, {"into": 999}, 404),
        (999, {"into": 1}, 404),
        (1, {}, 422),
        (1, {"into": 2, "force": True}, 422),
    ],
)
def test_bad_merges_are_rejected(
    duplicated: TestClient, source: int, body: dict[str, Any], expected: int
) -> None:
    response = duplicated.post(f"/applications/{source}/merge", headers=AUTH, json=body)
    assert response.status_code == expected
    assert len(_ids(duplicated)) == 4


def test_merging_needs_the_token(duplicated: TestClient) -> None:
    assert duplicated.post("/applications/1/merge", json={"into": 2}).status_code == 401
    assert len(_ids(duplicated)) == 4


def test_excel_export_fits_one_page_width_when_printed(seeded: TestClient) -> None:
    """Without this, printing or exporting to PDF drops the right-hand columns."""
    sheet = _sheet(seeded)
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.sheet_properties.pageSetUpPr.fitToPage
    assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 0)


def _history(tmp_path: Path, repository: Repository, **env: str) -> TestClient:
    """One application built from three emails: acknowledged, interviewed, rejected."""
    steps = [("other", "ack-1"), ("interview_invitation", "invite-2"), ("rejection", "no-3")]
    for days, (category, message_id) in enumerate(steps):
        raw = assessment(category, company="Nubaria", summary=f"{category} <b>summary</b>")
        repository.record(message(message_id, days=days), MessageAssessment.model_validate(raw))
    app = create_app(
        _settings(tmp_path, **env),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: T0 + timedelta(days=10),
    )
    return TestClient(app)


def test_each_application_lists_the_emails_behind_it(
    tmp_path: Path, repository: Repository
) -> None:
    with _history(tmp_path, repository) as client:
        page = client.get("/applications", headers=AUTH).text
        [row] = client.get("/applications?format=json", headers=AUTH).json()

    emails = re.search(r'<details class="emails">.*?</details>', page, re.S)
    assert emails
    assert "<summary>3 emails</summary>" in emails.group(0)
    listed = re.findall(r'<span class="category">([^<]*)</span>', emails.group(0))
    assert listed == ["rejection", "interview invitation", "other"], "newest first"
    assert "2026-07-03" in emails.group(0)
    assert "&lt;b&gt;summary&lt;/b&gt;" in emails.group(0)
    assert "<b>summary</b>" not in page

    assert [e["message_id"] for e in row["emails"]] == ["no-3", "invite-2", "ack-1"]
    assert row["emails"][0]["category"] == "rejection"


def test_fixture_emails_get_no_gmail_link(tmp_path: Path, repository: Repository) -> None:
    """Fixture ids are RFC Message-IDs; a Gmail link built from one would lead nowhere."""
    with _history(tmp_path, repository) as client:
        page = client.get("/applications", headers=AUTH).text
        [row] = client.get("/applications?format=json", headers=AUTH).json()
    assert "mail.google.com" not in page
    assert [e["url"] for e in row["emails"]] == [None, None, None]


def test_gmail_emails_link_to_the_original_message(tmp_path: Path, repository: Repository) -> None:
    with _history(tmp_path, repository, APPLY_AGENT_EMAIL_PROVIDER="gmail") as client:
        page = client.get("/applications", headers=AUTH).text
        [row] = client.get("/applications?format=json", headers=AUTH).json()

    assert row["emails"][0]["url"] == "https://mail.google.com/mail/u/0/#all/no-3"
    links = re.findall(r'<a href="([^"]*)" target="_blank" rel="noopener noreferrer">', page)
    assert links == [e["url"] for e in row["emails"]]


def test_a_message_id_cannot_break_out_of_the_link() -> None:
    assert (
        gmail_url('x"/><script>') == "https://mail.google.com/mail/u/0/#all/x%22%2F%3E%3Cscript%3E"
    )


def _offer_ids(client: TestClient) -> dict[str, dict[str, Any]]:
    rows = client.get("/offers?format=json", headers=AUTH).json()
    return {row["company"]: row for row in rows}


def test_offers_from_alerts_are_checked_against_applications(
    tmp_path: Path, repository: Repository
) -> None:
    repository.record(
        message("m1", "t1", days=1),
        MessageAssessment.model_validate(
            assessment(company="Corvina Retail", role="Senior ML Engineer")
        ),
    )
    repository.record(
        message("m2", "t2", days=1),
        MessageAssessment.model_validate(assessment(company="Brisa Mobility", role="Data Analyst")),
    )
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        clock=lambda: ALERT_DAY,
    )
    with TestClient(app) as client:
        assert "No offers yet" in client.get("/offers", headers=AUTH).text
        job = _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
        assert job["status"] == "succeeded"
        assert job["report"]["offers"]["new_offers"] == 3

        offers = _offer_ids(client)
        assert offers["Corvina Retail"]["section"] == "applied"
        assert offers["Corvina Retail"]["application_status"] == "applied"
        assert offers["Brisa Mobility"]["match"] == "same_company"
        assert offers["Brisa Mobility"]["section"] == "new"
        assert offers["Pendulum Health"]["match"] == "none"

        page = client.get("/offers", headers=AUTH).text
        assert "2 new" in page
        assert "You applied here for Data Analyst" in page
        assert "LinkedIn search is off" in page
        # The alert gave no link, so the page searches LinkedIn for the posting.
        assert "jobs/search/?keywords=AI+Engineer+Pendulum+Health&amp;location=" in page


def test_an_offer_can_be_marked_and_dismissed(tmp_path: Path, repository: Repository) -> None:
    app = create_app(_settings(tmp_path), run_sync=lambda: EMPTY_REPORT, repository=repository)
    with TestClient(app) as client:
        _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
        offer_id = _offer_ids(client)["Pendulum Health"]["id"]

        for state, section in (("interested", "interested"), ("dismissed", "dismissed")):
            response = client.put(f"/offers/{offer_id}/state", headers=AUTH, json={"state": state})
            assert response.json() == {"id": offer_id, "state": state}
            assert _offer_ids(client)["Pendulum Health"]["section"] == section

        bad = client.put(f"/offers/{offer_id}/state", headers=AUTH, json={"state": "applied"})
        assert bad.status_code == 422
        missing = client.put("/offers/999/state", headers=AUTH, json={"state": "new"})
        assert missing.status_code == 404


def test_offer_pages_need_the_token(client: TestClient) -> None:
    assert client.get("/offers").status_code == 401
    assert client.put("/offers/1/state", json={"state": "new"}).status_code == 401


def test_a_failed_offer_scan_does_not_fail_the_mail_sync(
    tmp_path: Path, repository: Repository
) -> None:
    def broken_scan() -> OfferScanReport:
        raise RuntimeError("inbox unreachable")

    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        run_offer_scan=broken_scan,
    )
    with TestClient(app) as client:
        job = _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
        assert job["status"] == "succeeded"
        assert job["report"]["processed"] == 0
        assert job["report"]["offers"] == {"error": "RuntimeError: inbox unreachable"}
        page = client.get("/offers", headers=AUTH).text
        assert "Looking for offers failed" in page
        assert 'onclick="sync(this)"' in page


def test_offers_are_ranked_by_fit_then_date_and_old_ones_folded(
    tmp_path: Path, repository: Repository
) -> None:
    offers = OfferRepository(repository.session_factory)

    def posting(title: str, linkedin_id: str, days_ago: int) -> JobOffer:
        return JobOffer(
            source=OfferSource.LINKEDIN_ALERT,
            linkedin_id=linkedin_id,
            title=title,
            company=f"Company {linkedin_id}",
            seen_at=ALERT_DAY - timedelta(days=days_ago),
        )

    offers.save(
        [
            posting("Senior AI Engineer", "4100000001", 0),
            posting("AI Engineer", "4100000002", 3),
            posting("Junior AI Engineer", "4100000003", 5),
            posting("AI Engineer", "4100000004", 1),
            posting("Junior Data Engineer", "4100000005", 30),
        ]
    )
    (tmp_path / "profile.toml").write_text('target = ["ai"]\n', encoding="utf-8")
    settings = _settings(tmp_path)
    app = create_app(
        settings,
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        offers=offers,
        clock=lambda: ALERT_DAY,
    )
    with TestClient(app) as client:
        rows = client.get("/offers?format=json", headers=AUTH).json()
        page = client.get("/offers", headers=AUTH).text

    assert [(r["linkedin_id"], r["fit"], r["section"]) for r in rows] == [
        ("4100000003", "good", "new"),  # entry title + keyword: 85
        ("4100000005", "good", "older"),  # entry title: 75, but no alert for 30 days
        ("4100000004", "possible", "new"),  # unstated + keyword: 65, newer
        ("4100000002", "possible", "new"),  # unstated + keyword: 65, older
        ("4100000001", "unlikely", "poor"),  # senior title for a new graduate
    ]
    assert rows[-1]["fit_reasons"] == [
        "A senior title, usually 5+ years, beyond your limit of 2",
        "Wanted: ai",
    ]
    assert "1 poor fit" in page
    assert "1 a good fit" in page
    assert "1 older than 14 days" in page
    assert page.index("Junior AI Engineer") < page.index("Senior AI Engineer")


def test_closed_offers_fold_away_and_teach_when_offers_go_old(
    tmp_path: Path, repository: Repository
) -> None:
    offers = OfferRepository(repository.session_factory)
    offers.save(
        [
            JobOffer(
                source=OfferSource.LINKEDIN_ALERT,
                linkedin_id=f"41000000{n:02d}",
                title=f"Role {n}",
                company=f"Company {n}",
                seen_at=ALERT_DAY - timedelta(days=n),
            )
            for n in range(1, 9)
        ]
    )
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        offers=offers,
        clock=lambda: ALERT_DAY,
    )
    with TestClient(app) as client:

        def by_id() -> dict[str, dict[str, Any]]:
            rows = client.get("/offers?format=json", headers=AUTH).json()
            return {row["linkedin_id"]: row for row in rows}

        page = client.get("/offers", headers=AUTH).text
        assert "Mark 5 more as <em>No longer available</em>" in page
        assert "No longer available</option>" in page
        assert by_id()["4100000008"]["section"] == "new"  # 8 days < the default 14

        # The user finds five offers closed, 2 to 6 days after their alert.
        for n in range(2, 7):
            offer_id = by_id()[f"41000000{n:02d}"]["id"]
            response = client.put(f"/offers/{offer_id}/closed", headers=AUTH)
            assert response.status_code == 200
            assert response.json()["closed_at"] is not None

        rows = by_id()
        assert {rows[f"41000000{n:02d}"]["section"] for n in range(2, 7)} == {"closed"}
        # Learned: offers close at about 4 days, so 7 and 8 days old now fold.
        assert [rows[f"41000000{n:02d}"]["section"] for n in (1, 7, 8)] == [
            "new",
            "older",
            "older",
        ]
        page = client.get("/offers", headers=AUTH).text
        assert "4 days after their last alert" in page
        assert "the typical age of the 5 offers you found closed" in page
        assert "5 no longer available" in page
        assert rows["4100000001"]["days_since_last_alert"] == 1

        # Undo one: four closures are too few to learn from, so 14 days again.
        three = rows["4100000003"]["id"]
        assert client.delete(f"/offers/{three}/closed", headers=AUTH).json()["closed_at"] is None
        rows = by_id()
        assert [rows[f"41000000{n:02d}"]["section"] for n in (3, 7, 8)] == ["new", "new", "new"]

        assert client.put("/offers/999/closed", headers=AUTH).status_code == 404
        assert client.delete("/offers/999/closed", headers=AUTH).status_code == 404


def _searcher(repository: Repository, *, infojobs: bool) -> JobSearcher:
    """A JobSearcher whose InfoJobs and job boards are fakes in their real shapes."""

    def handle(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "infojobs" in url:
            return httpx.Response(200, text=json.dumps({"items": [infojobs_item(1)]}))
        if "boards/arcwell" in url:
            return httpx.Response(200, text=json.dumps(GREENHOUSE))
        return httpx.Response(404)

    client = httpx.Client(transport=httpx.MockTransport(handle))
    sessions = repository.session_factory
    return JobSearcher(
        OfferRepository(sessions),
        SearchRepository(sessions),
        CandidateProfile(target=("ai",)),
        boards=CompanyBoards(client, clock=lambda: ALERT_DAY),
        infojobs=InfoJobsSearch(
            client, client_id="i", client_secret=INFOJOBS_SECRET, clock=lambda: ALERT_DAY
        )
        if infojobs
        else None,
        clock=lambda: ALERT_DAY,
    )


def _search_app(tmp_path: Path, repository: Repository, *, infojobs: bool) -> TestClient:
    return TestClient(
        create_app(
            _settings(tmp_path),
            run_sync=lambda: EMPTY_REPORT,
            repository=repository,
            run_offer_scan=lambda: OfferScanReport(),
            searcher=_searcher(repository, infojobs=infojobs),
            clock=lambda: ALERT_DAY,
        )
    )


def test_without_infojobs_the_page_says_how_to_turn_it_on(
    tmp_path: Path, repository: Repository
) -> None:
    with _search_app(tmp_path, repository, infojobs=False) as client:
        page = client.get("/search", headers=AUTH).text
        assert "APPLY_AGENT_INFOJOBS_CLIENT_ID" in page
        assert "<fieldset disabled>" in page
        assert "Companies you watch" in page
        searched = client.get("/search?what=python", headers=AUTH).text
        assert "No title search is set up for Spain" in searched


def test_a_search_shows_ranked_results_that_join_the_offers(
    tmp_path: Path, repository: Repository
) -> None:
    with _search_app(tmp_path, repository, infojobs=True) as client:
        page = client.get(
            "/search?what=python&where=Madrid&country=es&max_days_old=7", headers=AUTH
        ).text
        assert "1 found on InfoJobs, 1 new to your offers." in page
        assert "Desarrollador/a Backend Python" in page
        assert "Offers from InfoJobs" in page
        # The results can be reordered by fit or date, and poor fits hidden.
        assert '<option value="posted">Newest</option>' in page
        assert "Hide poor fits" in page
        assert 'data-posted="' in page
        # The same search on LinkedIn, in the place and country chosen.
        assert "jobs/search/?keywords=python&amp;location=Madrid%2C+Spain" in page
        other = client.get("/search?what=python&country=nl", headers=AUTH).text
        assert "jobs/search/?keywords=python&amp;location=Netherlands" in other
        [offer] = client.get("/offers?format=json", headers=AUTH).json()
        assert offer["source"] == "infojobs"
        assert offer["fit_reasons"][0].startswith("The description asks for 2+ years")

        assert client.get("/search?what=x&country=zz", headers=AUTH).status_code == 422


def test_saved_searches_run_on_every_sync(tmp_path: Path, repository: Repository) -> None:
    with _search_app(tmp_path, repository, infojobs=True) as client:
        saved = client.post("/search/saved", headers=AUTH, json={"what": "python", "country": "es"})
        assert saved.status_code == 201
        assert "python · Spain" in client.get("/search", headers=AUTH).text

        job = _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
        assert job["report"]["search"]["searches_run"] == 1
        assert job["report"]["search"]["new_offers"] == 1

        search_id = saved.json()["id"]
        assert client.delete(f"/search/saved/{search_id}", headers=AUTH).status_code == 200
        assert client.delete(f"/search/saved/{search_id}", headers=AUTH).status_code == 404


def test_a_watched_company_is_found_read_and_checked_on_sync(
    tmp_path: Path, repository: Repository
) -> None:
    with _search_app(tmp_path, repository, infojobs=False) as client:
        missing = client.post("/search/companies", headers=AUTH, json={"name": "Nowhere Ltd"})
        assert missing.status_code == 404
        assert "careers page" in missing.json()["detail"]

        request = {"name": "Arcwell", "keywords": "", "places": "Netherlands"}
        # Two postings under a guessed name: the user is asked first (D66).
        asked = client.post("/search/companies", headers=AUTH, json=request)
        assert asked.status_code == 409
        assert "greenhouse/arcwell with 2 posting(s)" in asked.json()["detail"]
        assert "AI Engineer" in asked.json()["detail"]

        watched = client.post("/search/companies", headers=AUTH, json=request | {"confirm": True})
        assert watched.status_code == 201
        assert watched.json() | {"id": 0} == {
            "id": 0,
            "board": "greenhouse",
            "slug": "arcwell",
            "offers": 1,  # the AI role; the sales role misses the profile's words
            "new": 1,
        }
        page = client.get("/search", headers=AUTH).text
        assert "Greenhouse" in page
        assert "Netherlands" in page

        job = _wait_for(client, client.post("/sync", headers=AUTH).json()["id"])
        assert job["report"]["search"]["companies_checked"] == 1

        company_id = watched.json()["id"]
        assert client.delete(f"/search/companies/{company_id}", headers=AUTH).status_code == 200
        assert len(client.get("/offers?format=json", headers=AUTH).json()) == 1


def test_the_offers_page_shows_offers_from_every_source(
    tmp_path: Path, repository: Repository
) -> None:
    """Regression: the page knew only LinkedIn's sources and failed on any other."""
    offers = OfferRepository(repository.session_factory)
    offers.save(
        [
            JobOffer(
                source=source,
                external_id=f"{source.value}/{n}",
                title=f"Role {n}",
                company=f"Company {n}",
                seen_at=ALERT_DAY,
            )
            for n, source in enumerate(OfferSource)
        ]
    )
    app = create_app(
        _settings(tmp_path),
        run_sync=lambda: EMPTY_REPORT,
        repository=repository,
        offers=offers,
        clock=lambda: ALERT_DAY,
    )
    with TestClient(app) as client:
        response = client.get("/offers", headers=AUTH)
    assert response.status_code == 200
    for label in ("Greenhouse", "Recruitee", "Personio", "SmartRecruiters", "InfoJobs"):
        assert f"{label}</span>" in response.text


def _company_app(
    tmp_path: Path,
    repository: Repository,
    summarizer: Any = None,
    llm: str = "baseline",
    investigator: Any = None,
) -> TestClient:
    def handle(request: httpx.Request) -> httpx.Response:
        if request.url.host == "news.google.com":
            return httpx.Response(200, text=GOOGLE_NEWS)
        if request.url.host == "hn.algolia.com":
            return httpx.Response(200, text=json.dumps({"hits": []}))
        if "boards/arcwell" in str(request.url):
            return httpx.Response(200, text=json.dumps(GREENHOUSE))
        return httpx.Response(404)

    client = httpx.Client(transport=httpx.MockTransport(handle))
    sessions = repository.session_factory
    searcher = JobSearcher(
        OfferRepository(sessions),
        SearchRepository(sessions),
        CandidateProfile(target=("ai",)),
        boards=CompanyBoards(client, clock=lambda: ALERT_DAY),
        clock=lambda: ALERT_DAY,
    )
    return TestClient(
        create_app(
            _settings(tmp_path, APPLY_AGENT_LLM=llm),
            run_sync=lambda: EMPTY_REPORT,
            repository=repository,
            run_offer_scan=lambda: OfferScanReport(),
            searcher=searcher,
            researcher=CompanyResearcher(client, clock=lambda: ALERT_DAY),
            summarizer=summarizer,
            investigator=investigator,
            clock=lambda: ALERT_DAY,
        )
    )


def test_a_company_page_shows_its_hiring_news_and_whom_to_contact(
    tmp_path: Path, repository: Repository
) -> None:
    with _company_app(tmp_path, repository) as client:
        client.post("/search/companies", headers=AUTH, json={"name": "Arcwell", "confirm": True})
        page = client.get("/company?name=Arcwell&context=robots", headers=AUTH).text

        assert "<h1>Arcwell</h1>" in page
        assert "<strong>1</strong> open" in page
        assert "AI Engineer" in page
        assert "Arcwell raises EUR 40M" in page
        assert "Careers page" in page
        assert "https://job-boards.greenhouse.io/arcwell" in page
        assert "linkedin.com/search/results/people/?keywords=Technical+Recruiter+Arcwell" in page
        assert "Set APPLY_AGENT_LLM to ollama or anthropic" in page
        offers = client.get("/offers", headers=AUTH).text
        assert 'href="company?name=Arcwell"' in offers

        assert client.get("/company?name=%20", headers=AUTH).status_code == 422
        missing = client.post("/company/summary", headers=AUTH, json={"name": "Arcwell"})
        assert missing.status_code == 409


def test_the_news_summary_comes_from_the_model(tmp_path: Path, repository: Repository) -> None:
    asked: list[tuple[str, int]] = []

    def summarizer(name: str, news: Sequence[NewsItem]) -> str:
        asked.append((name, len(news)))
        return "- It raised money."

    with _company_app(tmp_path, repository, summarizer) as client:
        page = client.get("/company?name=Arcwell", headers=AUTH).text
        assert "Summarize the news" in page
        response = client.post(
            "/company/summary", headers=AUTH, json={"name": "Arcwell", "context": None}
        )
        assert response.json() == {"summary": "- It raised money."}
        assert asked == [("Arcwell", 1)]


def test_a_list_of_companies_is_imported_checked_and_broken_down(
    tmp_path: Path, repository: Repository
) -> None:
    with _company_app(tmp_path, repository) as client:
        assert "No companies yet" in client.get("/companies", headers=AUTH).text
        started = client.post(
            "/companies/import", headers=AUTH, json={"text": "Netherlands: Arcwell y Nowhere"}
        )
        assert started.status_code == 202
        assert started.json()["added"] == 2
        job_id = started.json()["job"]["id"]
        deadline = time.monotonic() + 30
        while client.get(f"/companies/check/{job_id}", headers=AUTH).json()["status"] == "running":
            assert time.monotonic() < deadline
            time.sleep(0.05)

        page = client.get("/companies", headers=AUTH).text
        assert "<h2>Netherlands" in page
        assert "Maybe Greenhouse (arcwell)" in page  # two postings: the user decides
        assert "Not found by its name. Paste the link" in page  # Nowhere
        assert "1 to confirm" in page

        arcwell = next(
            t for t in TargetRepository(repository.session_factory).list() if t.name == "Arcwell"
        )
        confirmed = client.post(f"/companies/{arcwell.id}/confirm", headers=AUTH)
        assert confirmed.json() | {"id": 0} == {"id": 0, "offers": 1, "new": 1}
        page = client.get("/companies", headers=AUTH).text
        assert "<strong>1</strong> open" in page
        assert "AI Engineer" in page
        assert "job-boards.greenhouse.io/arcwell" in page
        assert client.post(f"/companies/{arcwell.id}/confirm", headers=AUTH).status_code == 409

        assert client.delete(f"/companies/{arcwell.id}", headers=AUTH).status_code == 200
        assert client.delete(f"/companies/{arcwell.id}", headers=AUTH).status_code == 404
        assert client.post("/companies/999/reject", headers=AUTH).status_code == 404

        # Not found by name: the user pastes its careers link (D71).
        nowhere = next(
            t for t in TargetRepository(repository.session_factory).list() if t.name == "Nowhere"
        )
        link = f"/companies/{nowhere.id}/board"
        refused = client.post(link, headers=AUTH, json={"url": "https://nowhere.example/jobs"})
        assert refused.status_code == 422
        assert "myworkdayjobs.com" in refused.json()["detail"]
        empty = client.post(
            link, headers=AUTH, json={"url": "https://job-boards.greenhouse.io/nobody"}
        )
        assert empty.status_code == 404
        linked = client.post(
            link, headers=AUTH, json={"url": "https://job-boards.greenhouse.io/arcwell"}
        )
        assert linked.json() == {"id": nowhere.id, "board": "greenhouse", "offers": 1, "new": 0}
        [found] = [
            t for t in TargetRepository(repository.session_factory).list() if t.name == "Nowhere"
        ]
        assert found.status is BoardStatus.FOUND
        assert client.delete(f"/companies/{nowhere.id}", headers=AUTH).status_code == 200


def test_without_searxng_the_company_page_says_how_to_research_online(
    tmp_path: Path, repository: Repository
) -> None:
    with _company_app(tmp_path, repository) as client:
        page = client.get("/company?name=Arcwell", headers=AUTH).text
        assert "Needs SearXNG" in page
        refused = client.post("/company/investigate", headers=AUTH, json={"name": "Arcwell"})
        assert refused.status_code == 409


def test_an_investigation_is_queued_run_and_shown_with_its_sources(
    tmp_path: Path, repository: Repository
) -> None:
    model = ScriptedChatModel(script=[AIMessage("- They hire junior ML engineers [1]")])
    pages = {"https://arcwell.example/careers": "<title>Careers</title>Junior ML roles."}
    research = web_investigator(model, pages)
    with _company_app(
        tmp_path, repository, summarizer=lambda name, news: "-", investigator=research
    ) as client:
        assert "Investigate</button>" in client.get("/company?name=Arcwell", headers=AUTH).text
        queued = client.post(
            "/company/investigate",
            headers=AUTH,
            json={"name": "Arcwell", "question": "Do they hire juniors?"},
        )
        assert queued.status_code == 202
        investigation_id = queued.json()["id"]
        deadline = time.monotonic() + 30
        state = {"status": "queued"}
        while state["status"] in ("queued", "running"):
            assert time.monotonic() < deadline
            time.sleep(0.05)
            state = client.get(f"/company/investigations/{investigation_id}", headers=AUTH).json()
        assert state["status"] == "done"

        page = client.get("/company?name=Arcwell", headers=AUTH).text
        assert ": Do they hire juniors?</h3>" in page
        assert "- They hire junior ML engineers [1]" in page
        assert "https://arcwell.example/careers" in page
        assert client.get("/company/investigations/999", headers=AUTH).status_code == 404


def test_a_companys_breakdown_says_where_and_under_which_names(
    tmp_path: Path, repository: Repository
) -> None:
    offers = OfferRepository(repository.session_factory)

    def alert(linkedin_id: str, title: str, company: str, location: str) -> JobOffer:
        return JobOffer(
            source=OfferSource.LINKEDIN_ALERT,
            linkedin_id=linkedin_id,
            title=title,
            company=company,
            location=location,
            seen_at=ALERT_DAY,
        )

    offers.save(
        [
            alert("4100000101", "Java Engineer Rail Signalling", "Arcwell", "Tres Cantos"),
            alert("4100000102", "AI Platform Engineer", "Arcwell", "Zug"),
            alert("4100000103", "AI Engineer", "Arcwell Energy", "Madrid"),
            alert("4100000104", "Data Engineer", "Arcwell", "Berlin, Germany"),
        ]
    )
    TargetRepository(repository.session_factory).add([("Alemania", "Arcwell")])
    with _company_app(tmp_path, repository) as client:
        page = client.get("/companies", headers=AUTH).text
        company = client.get("/company?name=Arcwell", headers=AUTH).text

    assert "2 in Spain" in page
    assert "<strong>1 in Germany</strong>" in page  # the group's own market
    assert "1 in Switzerland" in page
    assert "Includes Arcwell Energy (1)" in page
    assert 'Tres Cantos <span class="muted">(Spain)</span>' in company
    assert '<span class="quiet">Arcwell Energy</span>' in company


def test_noting_a_language_dismisses_the_offer_and_flags_the_company(
    tmp_path: Path, repository: Repository
) -> None:
    offers = OfferRepository(repository.session_factory)
    offers.save(
        [
            JobOffer(
                source=OfferSource.LINKEDIN_ALERT,
                linkedin_id=f"41000002{n:02d}",
                title=title,
                company="Arcwell",
                location="Utrecht",
                seen_at=ALERT_DAY,
            )
            for n, title in enumerate(["Junior Backend Engineer", "Junior Java Engineer"])
        ]
    )
    (tmp_path / "profile.toml").write_text(
        'target = ["backend"]\nlanguages = ["spanish", "english"]\n', encoding="utf-8"
    )
    with _company_app(tmp_path, repository) as client:
        page = client.get("/offers", headers=AUTH).text
        assert 'data-language="Dutch"' in page  # suggested from Utrecht
        first, second = sorted(
            client.get("/offers?format=json", headers=AUTH).json(), key=lambda o: o["title"]
        )
        assert second["fit"] == "good"

        noted = client.post(
            f"/offers/{second['id']}/language", headers=AUTH, json={"language": "Nederlands"}
        )
        assert noted.json() == {"id": second["id"], "company": "Arcwell", "language": "dutch"}

        rows = {o["id"]: o for o in client.get("/offers?format=json", headers=AUTH).json()}
        assert rows[second["id"]]["section"] == "dismissed"
        assert "This company asked for Dutch before" in rows[first["id"]]["fit_reasons"]
        assert rows[first["id"]]["fit"] == "unlikely"

        bad = client.post(f"/offers/{first['id']}/language", headers=AUTH, json={"language": "x?"})
        assert bad.status_code == 422
        gone = client.post("/offers/999/language", headers=AUTH, json={"language": "dutch"})
        assert gone.status_code == 404
